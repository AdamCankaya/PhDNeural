"""Generate Bio-NAS broader-impacts white paper DOCX from papers/*.md source.

Requires: pip install python-docx

Usage (from repo root or Bio-NAS):
  python Bio-NAS/scripts/generate_ethics_white_paper_docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

BIO_NAS_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MD = BIO_NAS_ROOT / "papers" / "white_paper_bio_nas_ethics_broader_impacts.md"
DEFAULT_DOCX = BIO_NAS_ROOT / "papers" / "white_paper_bio_nas_ethics_broader_impacts.docx"


def set_run_font(run, name="Times New Roman", size=10, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def style_paragraph(p, space_before=0, space_after=6, line_spacing=1.08, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if align is not None:
        p.alignment = align


def add_centered(doc, text, size=10, bold=False, italic=False, space_before=0, space_after=6):
    p = doc.add_paragraph()
    style_paragraph(
        p, space_before=space_before, space_after=space_after, align=WD_ALIGN_PARAGRAPH.CENTER
    )
    run = p.add_run(strip_md_inline(text))
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_body_with_inline(doc, text, size=10, italic=False, indent=False):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=0, space_after=6)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.2)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        chunk = part[2:-2] if bold else part
        chunk = strip_md_inline(chunk)
        run = p.add_run(chunk)
        set_run_font(run, size=size, bold=bold, italic=italic and not bold)
    return p


def strip_md_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("*", "")
    return text


def add_heading_ieee(doc, text, level=1):
    p = doc.add_paragraph()
    clean = strip_md_inline(text)
    if level == 1:
        style_paragraph(p, space_before=12, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(clean.upper())
        set_run_font(run, size=10, bold=True)
    elif level == 2:
        style_paragraph(p, space_before=10, space_after=4)
        run = p.add_run(clean)
        set_run_font(run, size=10, bold=True)
    elif level == 3:
        style_paragraph(p, space_before=8, space_after=3)
        run = p.add_run(clean)
        set_run_font(run, size=10, bold=True, italic=True)
    else:
        style_paragraph(p, space_before=6, space_after=3)
        run = p.add_run(clean)
        set_run_font(run, size=10, italic=True)
    return p


def add_table_from_markdown(doc, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = strip_md_inline(cell_text.strip())
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=9, bold=(r_idx == 0))


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def build_document_from_markdown(md_path: Path, out_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)

    i = 0
    title_done = False
    table_buffer: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_table_from_markdown(doc, table_buffer)
            table_buffer = []
            doc.add_paragraph()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|"):
            if not stripped.replace("|", "").replace("-", "").strip():
                i += 1
                continue
            table_buffer.append(parse_table_row(line))
            i += 1
            continue
        flush_table()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("## "):
            title = stripped[2:].strip()
            add_centered(doc, title, size=14, bold=True, space_after=8)
            title_done = True
            i += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            subtitle = stripped.strip("*").strip()
            add_centered(doc, subtitle, size=10, italic=True, space_after=10)
            i += 1
            continue

        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            if heading.lower() == "authoring note":
                add_heading_ieee(doc, "Authoring Note", level=2)
                i += 1
                note_lines: list[str] = []
                while i < len(lines) and not lines[i].strip().startswith("## "):
                    if lines[i].strip():
                        note_lines.append(lines[i].strip())
                    i += 1
                note_text = " ".join(note_lines)
                add_body_with_inline(doc, note_text, italic=True)
                continue
            add_heading_ieee(doc, heading, level=1)
            i += 1
            continue

        if stripped.startswith("### "):
            add_heading_ieee(doc, stripped[4:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith("**") and ":**" in stripped:
            add_heading_ieee(doc, stripped.rstrip(":"), level=3)
            i += 1
            continue

        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith("#") or nxt.startswith("|") or (
                nxt.startswith("*") and nxt.endswith("*") and not nxt.startswith("**")
            ):
                break
            if nxt.startswith("**") and nxt.endswith(":**"):
                break
            if nxt.startswith("### "):
                break
            para_lines.append(nxt)
            i += 1
        add_body_with_inline(doc, " ".join(para_lines), indent=title_done)

    flush_table()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> int:
    md_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_MD
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_DOCX
    if not md_path.is_file():
        print(f"Missing markdown source: {md_path}", file=sys.stderr)
        return 1
    build_document_from_markdown(md_path, out_path)
    print(out_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
